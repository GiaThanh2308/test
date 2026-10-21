"""Nhập danh sách lớp từ Excel để sửa TÊN CÓ DẤU (và cập nhật LỚP khi xáo lớp).

File không cần mã học sinh: hệ thống ghép theo (tên bỏ dấu + lớp). Nhãn khuôn mặt (face_label) và mã HS không bị đụng tới.
Hỗ trợ: một sheet có cột Lớp, hoặc mỗi lớp một sheet (tên sheet = tên lớp); cột "Họ và tên" hoặc tách "Họ đệm" + "Tên".
Phần tính toán ở đây không phụ thuộc cơ sở dữ liệu nên kiểm thử được.
"""
import re
import unicodedata
from collections import defaultdict
from io import BytesIO
from typing import Optional

from openpyxl import load_workbook


def plain(text: Optional[str]) -> str:
    """Bỏ dấu (kể cả đ), chữ thường, gom khoảng trắng."""
    nfkd = unicodedata.normalize("NFKD", (text or "").replace("đ", "d").replace("Đ", "D"))
    s = "".join(c for c in nfkd if not unicodedata.combining(c)).lower()
    return re.sub(r"\s+", " ", s).strip()


def clean_name(text) -> str:
    return re.sub(r"\s+", " ", unicodedata.normalize("NFC", str(text or ""))).strip()


def class_key(text) -> str:
    return re.sub(r"\s+", "", str(text or "")).upper()


def _classify_header(cell) -> Optional[str]:
    h = plain(str(cell)) if cell is not None else ""
    if not h:
        return None
    if "lop" in h.split() or h in ("lop", "lop hoc"):
        return "class"
    if h in ("ho", "ho dem", "ho lot", "ho va ten dem", "ho va ten lot", "ho ten dem"):
        return "ho"
    if h == "ten":
        return "ten"
    if "ho" in h.split() and "ten" in h.split():
        return "full"
    if h in ("ten hoc sinh", "hoc sinh", "ho ten", "hoten"):
        return "full"
    return None


def parse_roster(data: bytes) -> tuple[list[dict], list[str]]:
    """→ (rows[{name, class_name, where}], cảnh báo). Ném ValueError nếu không nhận ra cột."""
    wb = load_workbook(BytesIO(data), data_only=True, read_only=True)
    rows, warnings = [], []
    for ws in wb.worksheets:
        values = list(ws.iter_rows(values_only=True, max_row=5000))
        header_idx, cols = None, {}
        for i, row in enumerate(values[:15]):
            found = {}
            for j, cell in enumerate(row):
                kind = _classify_header(cell)
                if kind and kind not in found:
                    found[kind] = j
            if "full" in found or ("ten" in found and "ho" in found):
                header_idx, cols = i, found
                break
        if header_idx is None:
            warnings.append(f"Sheet '{ws.title}': không thấy dòng tiêu đề có cột Họ và tên — bỏ qua")
            continue
        for i, row in enumerate(values[header_idx + 1:], start=header_idx + 2):
            def cell(kind):
                j = cols.get(kind)
                return row[j] if j is not None and j < len(row) else None
            if "full" in cols:
                name = clean_name(cell("full"))
            else:
                name = clean_name(f"{cell('ho') or ''} {cell('ten') or ''}")
            if not name or name.isdigit() or plain(name) in ("stt", "ho va ten"):
                continue
            cls = class_key(cell("class")) if "class" in cols else class_key(ws.title)
            if not cls:
                warnings.append(f"Sheet '{ws.title}' dòng {i}: thiếu lớp — bỏ qua")
                continue
            rows.append({"name": name, "class_name": cls, "where": f"{ws.title}!{i}"})
    if not rows:
        raise ValueError("Không đọc được học sinh nào. File cần dòng tiêu đề có 'Họ và tên' (hoặc 'Họ đệm' + 'Tên') "
                         "và cột 'Lớp' (hoặc mỗi lớp một sheet, tên sheet là tên lớp).")
    return rows, warnings


def plan_import(rows: list[dict], students: list[dict]) -> dict:
    """students: [{id, full_name, class_name}] (chỉ học sinh đang học).
    Trả về kế hoạch; không sửa gì."""
    by_nc, by_name = defaultdict(list), defaultdict(list)
    for s in students:
        by_nc[(plain(s["full_name"]), class_key(s["class_name"]))].append(s)
        by_name[plain(s["full_name"])].append(s)

    file_count = defaultdict(int)
    for r in rows:
        file_count[(plain(r["name"]), class_key(r["class_name"]))] += 1

    updates, ambiguous, not_found, matched_ids = [], [], [], set()
    pending_moves = []
    unchanged = 0

    for r in rows:
        nk, ck = plain(r["name"]), class_key(r["class_name"])
        cands = by_nc.get((nk, ck), [])
        if file_count[(nk, ck)] > 1 or len(cands) > 1:
            ambiguous.append({"name": r["name"], "class_name": r["class_name"], "where": r["where"],
                              "reason": "Trùng tên trong cùng lớp (file hoặc hệ thống) — sửa tay"})
            matched_ids.update(c["id"] for c in cands)
            continue
        if len(cands) == 1:
            s = cands[0]
            matched_ids.add(s["id"])
            if s["full_name"] == r["name"]:
                unchanged += 1
            else:
                updates.append({"id": s["id"], "old_name": s["full_name"], "new_name": r["name"],
                                "old_class": s["class_name"], "new_class": s["class_name"], "move": False})
            continue
        pending_moves.append(r)

    for r in pending_moves:  # không có trong lớp đó → thử tìm theo tên trong cả trường (xáo lớp)
        nk = plain(r["name"])
        cands = [s for s in by_name.get(nk, []) if s["id"] not in matched_ids]
        if len(cands) == 1 and len([x for x in pending_moves if plain(x["name"]) == nk]) == 1:
            s = cands[0]
            matched_ids.add(s["id"])
            updates.append({"id": s["id"], "old_name": s["full_name"], "new_name": r["name"],
                            "old_class": s["class_name"], "new_class": r["class_name"], "move": True})
        elif cands:
            ambiguous.append({"name": r["name"], "class_name": r["class_name"], "where": r["where"],
                              "reason": "Không có trong lớp này, nhưng có nhiều người trùng tên ở lớp khác"})
        else:
            not_found.append({"name": r["name"], "class_name": r["class_name"], "where": r["where"]})

    file_classes = {class_key(r["class_name"]) for r in rows}
    not_in_list = [{"id": s["id"], "name": s["full_name"], "class_name": s["class_name"]}
                   for s in students if s["id"] not in matched_ids and class_key(s["class_name"]) in file_classes]

    return {"updates": updates, "unchanged": unchanged, "ambiguous": ambiguous,
            "not_found": not_found, "not_in_list": not_in_list,
            "name_changes": sum(1 for u in updates if u["old_name"] != u["new_name"]),
            "moves": sum(1 for u in updates if u["move"])}
