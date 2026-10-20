"""Xuất Excel vi phạm theo mẫu sổ thi đua (kiểu đầu trang: nền xanh đậm 17365D, chữ trắng, viền vừa).

Sheet 1  CHI TIET      : THỜI GIAN · HỌ TÊN · LỚP · LỖI VI PHẠM (+ ghi chú, người ghi)
Sheet 2  NHAP VI PHAM  : cùng cột với sheet NHAP VI PHAM của sổ thi đua (TUẦN, LỚP, MÃ LỖI, SỐ LẦN, ...)
                         → copy cột A:D dán thẳng vào sổ.
"""
from collections import Counter
from datetime import date, datetime, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from backend.violation_types import CODE_POINTS, TYPE_TO_CODE

FONT = "Calibri"
HEADER_FILL = PatternFill("solid", fgColor="17365D")
HEADER_FONT = Font(name=FONT, size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name=FONT, size=11)
SIDE = Side(style="medium")
BORDER = Border(left=SIDE, right=SIDE, top=SIDE, bottom=SIDE)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def school_week(local_dt: datetime, week1_start: date) -> int:
    """Tuần học thứ mấy (tuần 1 bắt đầu từ ngày week1_start, thường là thứ Hai). <1 → 0 (chưa vào năm học)."""
    return max(0, (local_dt.date() - week1_start).days // 7 + 1)


def _style_sheet(ws, headers: list[str], widths: list[float]):
    ws.append(headers)
    ws.row_dimensions[1].height = 33
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font, c.fill, c.border, c.alignment = HEADER_FONT, HEADER_FILL, BORDER, CENTER
    ws.freeze_panes = "A2"


def _style_body(ws, left_cols: set[int]):
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font, c.border = BODY_FONT, BORDER
            c.alignment = LEFT if c.column in left_cols else CENTER


def build_workbook(rows: list[dict], week1_start: date) -> bytes:
    """rows: dict(time=datetime giờ địa phương, name, class_name, type, note, by), đã sắp theo thời gian."""
    wb = Workbook()

    ws = wb.active
    ws.title = "CHI TIET"
    _style_sheet(ws, ["STT", "THỜI GIAN", "HỌ VÀ TÊN", "LỚP", "LỖI VI PHẠM", "GHI CHÚ", "NGƯỜI GHI"],
                 [7, 19, 28, 10, 28, 32, 16])
    for i, r in enumerate(rows, 1):
        ws.append([i, r["time"].strftime("%d/%m/%Y %H:%M"), r["name"], r["class_name"],
                   r["type"], r["note"] or "", r["by"] or ""])
    _style_body(ws, left_cols={3, 5, 6})
    ws.auto_filter.ref = f"A1:G{max(ws.max_row, 2)}"

    ws2 = wb.create_sheet("NHAP VI PHAM")
    _style_sheet(ws2, ["TUẦN", "LỚP", "MÃ LỖI", "SỐ LẦN", "ĐIỂM QUY ĐỊNH", "ĐIỂM ÁP DỤNG 20–30",
                       "TỔNG ĐIỂM TRỪ", "GHI CHÚ / CĂN CỨ"], [12, 17, 15, 13, 19, 25, 20, 48])
    counts: Counter = Counter()
    unmapped: Counter = Counter()
    for r in rows:
        week, code = school_week(r["time"], week1_start), TYPE_TO_CODE.get(r["type"])
        if code:
            counts[(week, r["class_name"], code, r["type"])] += 1
        else:
            unmapped[(week, r["class_name"], r["type"])] += 1
    for (week, cls, code, type_name), n in sorted(counts.items()):
        pts = CODE_POINTS.get(code)
        ws2.append([week, cls, code, n, pts if pts is not None else "", "",
                    pts * n if pts is not None else "", type_name])
    for (week, cls, type_name), n in sorted(unmapped.items()):
        ws2.append([week, cls, "", n, "", "", "", f"{type_name} — chưa có mã lỗi trong sổ, tự gán mã"])
    _style_body(ws2, left_cols={8})

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def default_range(today: date) -> tuple[date, date]:
    start = today - timedelta(days=today.weekday())
    return start, start + timedelta(days=6)
